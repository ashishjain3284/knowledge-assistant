#!/usr/bin/env python
"""
generate_documents.py
---------------------
Creates the eight fictional company policy documents in the documents/ folder.

They are committed to the repository already, so you only need this script if
you want to regenerate or edit them.

    python scripts/generate_documents.py

.md files are written directly. .docx needs python-docx and .pdf needs
reportlab, both of which are in requirements.txt.
"""

from __future__ import annotations

import sys
from pathlib import Path

DOCS = Path(__file__).resolve().parents[1] / "documents"

# --------------------------------------------------------------------------- #
# The documents. Each is (filename, title, body). Everything is fictional.
# --------------------------------------------------------------------------- #

LEAVE_POLICY = """
1. Purpose and scope

This policy sets out the paid and unpaid leave available to all permanent
employees of Nimbus Technologies. It applies from the first day of employment
unless stated otherwise. Contractors and agency staff are covered by their own
agreements and are out of scope.

2. Annual leave entitlement

Full-time employees receive 25 days of paid annual leave per calendar year, in
addition to public holidays. Part-time employees receive a pro-rata entitlement
calculated on contracted hours. Employees who join part way through the year
accrue 2.08 days for each complete month of service.

After five years of continuous service, entitlement increases to 27 days. After
ten years it increases to 30 days. The increase takes effect from the January
following the service anniversary.

3. Carry-forward of unused leave

Employees may carry forward a maximum of 5 unused annual leave days into the
following calendar year. Carried-forward days must be taken by 31 March of that
year, after which they lapse and are not paid out.

Carry-forward above 5 days is permitted only where the employee was prevented
from taking leave by business demands, and requires written approval from both
the line manager and the HR Business Partner before 31 December.

Unused leave is not paid in lieu except on termination of employment, where it
is paid at the basic daily rate.

4. Booking annual leave

Leave is requested through the HR Portal and must be approved by the line
manager before travel is booked. The required notice is twice the length of the
leave requested, with a minimum of five working days. Requests of ten
consecutive days or more require four weeks' notice.

A manager may decline a request where team cover cannot be arranged. A declined
request must be answered within five working days with a reason.

No more than 40 per cent of a team may be on leave at the same time. During the
financial year-end period, 15 March to 5 April, leave in the Finance department
is granted only in exceptional circumstances.

5. Sick leave

Employees receive up to 10 days of paid sick leave per calendar year. Absence of
three consecutive days or more requires a medical certificate. Absence must be
reported to the line manager before 10:00 on the first day, by telephone rather
than message where possible.

Sick leave does not accrue and cannot be carried forward. Long-term absence
beyond 10 days moves to the company sick pay scheme, which pays full salary for
a further 8 weeks followed by statutory pay.

6. Parental leave

Primary caregivers receive 26 weeks of paid parental leave. Secondary caregivers
receive 8 weeks of paid parental leave, which may be taken in up to three blocks
within the first year. Both are available after 26 weeks of continuous service.

Adoption leave is granted on the same terms as parental leave. Employees
returning from parental leave may request a phased return of up to eight weeks
at reduced hours on full pay.

7. Other leave

Compassionate leave of up to 5 paid days is available on the death of an
immediate family member. Marriage leave is 3 paid days, available once during
employment. Examination leave of up to 5 days per year is available for
company-approved study.

Unpaid leave of up to three months may be requested after two years of service
and requires director approval. Benefits continue during unpaid leave but
annual leave does not accrue.

8. Public holidays

Employees receive the public holidays of their work location. Where an employee
is required to work a public holiday, they receive a day in lieu, to be taken
within three months.
"""

EMPLOYEE_HANDBOOK = """
1. Welcome

This handbook describes how we work at Nimbus Technologies. It sits alongside
your contract of employment. Where the two differ, the contract takes
precedence. The handbook is reviewed annually by the People team.

2. Working hours

Standard hours are 37.5 per week, normally 09:00 to 17:30 with an unpaid hour
for lunch. Core hours, when all employees are expected to be contactable, are
10:00 to 16:00 local time. Outside core hours, start and finish times are
flexible by agreement with the line manager.

Overtime is not routinely paid. Employees below manager grade who work more than
two additional hours in a day at the company's request may take time off in
lieu, agreed in advance.

3. Probation

New employees serve a probationary period of six months. During probation the
notice period is two weeks on either side. A probation review is held at three
months and again at six months. Probation may be extended once, by up to three
months, with written reasons.

4. Performance and development

Formal performance reviews are held twice a year, in April and October. Each
employee agrees objectives with their manager at the start of the cycle.
Reviews are recorded in the HR Portal and inform the annual salary review, which
takes effect on 1 July.

Every employee has an annual learning budget; see the Benefits Guide for the
current amount and how to claim it.

5. Salary and payment

Salaries are paid monthly on the 25th of the month, or the preceding working day
where the 25th falls on a weekend or public holiday. Payslips are available in
the HR Portal from two working days before payday.

6. Grievance procedure

An employee with a concern should first raise it informally with their line
manager. Where that is not appropriate or does not resolve the matter, a formal
grievance may be raised in writing with the HR Business Partner.

A formal grievance is acknowledged within three working days and a hearing held
within ten working days. The employee may be accompanied by a colleague. The
outcome is given in writing within five working days of the hearing, with a
right of appeal within ten working days.

7. Disciplinary procedure

The disciplinary process has four stages: informal discussion, written warning,
final written warning, and dismissal. Warnings remain live for twelve months.
Gross misconduct may result in dismissal without notice; examples are given in
the Code of Conduct.

8. Leaving the company

The notice period after probation is one month for employees below manager
grade, and three months for managers and above. Notice must be given in writing.

An exit interview is offered to every leaver and is conducted by a member of the
People team who is not in the leaver's reporting line. All company equipment
must be returned on or before the last working day. Final pay, including any
untaken annual leave, is paid in the payroll run following the leaving date.

9. Dress code

Dress is business casual. Client-facing meetings may require business dress; the
account lead will confirm. There is no dress code on days when an employee works
from home and has no external meetings.

10. References

The company provides factual references only, confirming job title and dates of
employment. Requests should be sent to the People team and not answered
personally by managers.
"""

BENEFITS_GUIDE = """
1. Overview

This guide summarises the benefits available to permanent employees of Nimbus
Technologies. Benefits start on the first day of employment unless a qualifying
period is stated. The company reviews the benefits package each year in June.

2. Private medical insurance

All employees receive private medical insurance for themselves at no cost.
Cover may be extended to a partner and dependent children at the employee's
expense through payroll deduction, at a preferential group rate.

The scheme covers inpatient and day-patient treatment in full, and outpatient
treatment up to 1,500 currency units per policy year. Pre-existing conditions
are excluded for the first two years of cover.

3. Dental and optical

Dental cover reimburses 80 per cent of routine treatment up to 400 units per
year. An eye test is reimbursed annually, with a contribution of up to 150 units
towards glasses where they are required for display screen work.

4. Life assurance and income protection

Life assurance is provided at four times basic annual salary, payable to
nominated beneficiaries. Employees should keep their nomination up to date in
the HR Portal.

Income protection pays 60 per cent of basic salary after 26 weeks of continuous
absence, until return to work or retirement.

5. Retirement savings

The company matches employee pension contributions up to 8 per cent of basic
salary. Employees are enrolled automatically at 5 per cent after three months of
service and may opt out or vary their contribution at any time.

6. Wellness allowance

Each employee receives a wellness allowance of 500 units per calendar year. It
may be spent on gym membership, fitness classes, mental health support,
ergonomic equipment for home use, or a health assessment. Claims are made
through the expenses system with a receipt, and the allowance does not carry
forward.

7. Learning and development budget

Each employee has a learning budget of 1,200 units per calendar year for
courses, conferences, certifications and books relevant to their role. Spending
above 600 units in a single item requires line manager approval in advance.

Where the company funds a qualification costing more than 3,000 units, the
employee agrees to repay a proportion if they leave within two years.

8. Employee share purchase plan

Employees with more than one year of service may contribute up to 10 per cent of
salary to the share purchase plan. Shares are purchased twice a year at a 15 per
cent discount to the market price.

9. Referral bonus

An employee who refers a successful candidate receives a bonus of 2,000 units,
paid after the new joiner completes probation. The referral must be logged in
the HR Portal before the candidate applies.

10. Other benefits

Employees receive one paid volunteering day per year, a season ticket loan for
commuting, and access to a discount platform for retail and travel.
"""

IT_POLICY = """
1. Purpose

This policy governs the use of Nimbus Technologies information systems,
devices, networks and data. It applies to all employees, contractors and anyone
given access to company systems. Acceptance is a condition of that access.

2. Accounts and authentication

Each person has a single named account. Accounts must not be shared, and one
person must never work under another person's credentials.

Passwords must be at least 12 characters and are changed every 90 days. The last
ten passwords cannot be reused. Multi-factor authentication is mandatory for
email, the VPN, the HR Portal and any administrative access.

Accounts lock after five failed sign-in attempts and unlock automatically after
30 minutes. The Service Desk will never ask for a password, and will never reset
one over chat or email.

3. Company devices

Laptops remain company property and must be returned on leaving. Devices must
not be modified, and the disk encryption and endpoint protection agents must not
be disabled.

Devices must be locked whenever unattended and must be restarted at least weekly
so that security updates apply. Loss or theft must be reported to the Service
Desk within 24 hours so the device can be wiped remotely.

4. Software

Software may only be installed from the Company Portal. Installing software from
the internet is blocked and attempts are logged. Additional licensed software is
requested through a ticket, with a business justification and a cost centre.

Use of unlicensed software, or of personal licences for company work, is not
permitted.

5. Data classification and handling

Information is classified as Public, Internal, Confidential or Restricted.
Confidential and Restricted information must not be copied to personal devices,
personal cloud storage or personal email under any circumstances.

Restricted data, which includes payroll, health and customer payment
information, may only be accessed on a company device over the VPN, and only by
those with an approved business need.

Removable media is blocked by default. Where a business need exists, an
encrypted device is issued by IT for a fixed period.

6. Acceptable personal use

Reasonable personal use of company systems is permitted where it does not
interfere with work, incur cost, or breach any policy. Personal use should be
occasional and should not include streaming, gaming, cryptocurrency activity, or
storing large personal files.

7. Email and messaging

External email is scanned and suspicious messages are quarantined. Employees
must report a suspected phishing message using the Report Phishing button rather
than deleting it.

Company information must not be forwarded to a personal address, including for
convenience when travelling.

8. Use of artificial intelligence tools

Only company-approved AI tools may be used with company information.
Confidential or Restricted information must never be entered into a public AI
service. Output from an AI tool must be reviewed by a person before it is used
in a customer-facing document or a decision.

9. Monitoring

The company monitors systems for security and compliance. Monitoring is
proportionate, logged, and covers system and network activity rather than the
content of personal communications, except where there is a specific and
authorised investigation.

10. Incident reporting

Any suspected security incident, data loss or unauthorised access must be
reported to the Service Desk immediately and in any case within 24 hours.
Reporting a genuine mistake promptly will not lead to disciplinary action; not
reporting it may.

11. Bring your own device

Personal phones may access company email and calendar through the managed
application only, with a device passcode and remote wipe of company data
enabled. Personal laptops must not be used for company work.
"""

TRAVEL_POLICY = """
1. Purpose and principles

This policy covers business travel and expenses for all employees of Nimbus
Technologies. Travel should be necessary, booked in advance, and at reasonable
cost. Where a video meeting will achieve the same outcome, it should be
preferred.

2. Approval

All travel requires line manager approval before booking. Trips with a total
cost above 2,000 units, and all international travel, additionally require
director approval.

Approval is requested in the travel system, which records the purpose, the
dates, the estimated cost and the approver.

3. Booking

Travel must be booked through the company travel platform. Bookings made outside
the platform are reimbursed only where the platform was unavailable and the
manager has confirmed this in writing.

Flights and hotels should be booked at least fourteen days in advance where the
trip is known. Late bookings require a note explaining why.

4. Flights

Economy class is standard for all flights. Premium economy may be booked where
the scheduled flight time exceeds six hours. Business class requires director
approval and is normally limited to flights over nine hours where the traveller
works on the day of arrival.

5. Accommodation

Hotel limits per night, excluding tax, are 180 units in London, Tokyo and New
York; 130 units in other major cities; and 100 units elsewhere. Where a
conference hotel exceeds the limit, the conference rate may be used with manager
approval.

6. Meals and daily allowance

A daily allowance covers meals and incidental costs and does not require
receipts. The allowance is 55 units per full day of international travel and 35
units per full day of domestic travel. A partial day of eight hours or less is
paid at half rate.

Where meals are provided by a conference or a host, the allowance is reduced by
40 per cent for that day. Alcohol is not reimbursable except at a client dinner
approved in advance by a director.

7. Ground transport

Public transport should be used where practical. Taxis and ride-hailing are
reimbursable for airport transfers, for journeys with luggage, and for travel
after 21:00 or before 06:00.

Mileage in a private vehicle is reimbursed at 0.45 units per kilometre for the
first 10,000 kilometres in a year and 0.25 units per kilometre thereafter. A
journey log must be kept. Parking and tolls are reimbursable; fines are not.

Car hire is limited to compact class unless three or more people are travelling
together.

8. Submitting expenses

Expenses are submitted in the expenses system within 30 days of the date of
expenditure. Claims submitted after 60 days are not reimbursed except in
exceptional circumstances approved by Finance.

An itemised receipt is required for every item above 25 units. A card statement
is not an acceptable substitute for a receipt.

Approved claims are paid in the next payroll run, or within ten working days for
claims above 1,000 units where the employee requests earlier payment.

9. Non-reimbursable items

The following are not reimbursed: fines and penalties, personal entertainment,
in-flight purchases other than meals on long-haul flights, hotel minibar,
upgrades not approved in advance, travel insurance purchased separately as the
company provides cover, and expenses for a partner or family member accompanying
the traveller.

10. Travel safety

The company provides travel insurance for all business travel; the policy number
is available in the travel system. Travellers must register their itinerary in
the travel system so they can be contacted in an emergency. Travel to countries
with a government advisory against travel requires director approval and a risk
assessment.
"""

REMOTE_WORK_POLICY = """
1. Purpose

This policy sets out how employees of Nimbus Technologies may work away from a
company office. It applies to all permanent employees who have completed
probation.

2. Hybrid working standard

The company operates a hybrid model. Employees are expected in a company office
for a minimum of three days per week, of which Tuesday is a fixed company-wide
in-office day for all teams.

Teams may agree a second fixed day. Individual arrangements that differ from the
three-day standard require director approval and are reviewed every six months.

3. Eligibility

Remote working is available after probation. Roles requiring physical presence,
including facilities, reception and hardware support, are not eligible and this
is stated in the job description.

An employee whose performance is under formal review may have remote working
paused for the duration of the review.

4. Home working environment

Employees are responsible for a safe and suitable working environment. A display
screen equipment self-assessment must be completed annually in the HR Portal.

The company provides a home office allowance of 400 units on joining, and a
further 200 units every three years, for a desk, chair, monitor or other
equipment. Items purchased with the allowance remain company property above a
value of 250 units.

A laptop, monitor, keyboard, mouse and headset are provided by IT on request and
are not deducted from the allowance.

5. Availability and communication

Employees working remotely are expected to be contactable during core hours,
10:00 to 16:00 local time, and to keep their calendar accurate. Cameras are
expected to be on for meetings with clients and for one-to-one meetings, and are
optional otherwise.

6. Security when working remotely

The VPN must be used for any access to internal systems. Work must not be done
on public Wi-Fi without the VPN. Screens must not be visible to others in a
public place when Confidential or Restricted information is displayed.

Company devices must not be used by family members. Paper documents containing
company information must be stored securely and shredded rather than placed in
household waste.

7. Working from another country

Working from another country is permitted for up to 20 working days per calendar
year, with line manager approval and a note in the HR Portal, and only from
countries on the approved list maintained by the People team.

Periods beyond 20 days create tax and employment obligations and require
approval from the People Director and Finance before travel is booked.

8. Expenses while working remotely

Home internet, heating and electricity are not reimbursed, as the home office
allowance covers the cost of working from home. Travel between home and a
company office is commuting and is not reimbursable, including on days when an
employee is asked to attend at short notice.
"""

CODE_OF_CONDUCT = """
# Code of Conduct

## 1. Our standards

Everyone at Nimbus Technologies is expected to act with honesty, respect and
care. This code sets the minimum standard. Where a local law is stricter, the
law applies.

Managers carry additional responsibility: to model the standards, to make it
safe to raise concerns, and to act on what they are told.

## 2. Respect at work

We do not tolerate harassment, bullying, discrimination or victimisation on any
ground, including age, disability, gender, gender identity, race, religion,
sexual orientation, or family status.

Harassment includes conduct that a reasonable person would find offensive, even
where no offence was intended, and includes conduct online and in messaging
tools.

## 3. Conflicts of interest

A conflict of interest arises where a personal interest could influence, or
appear to influence, a business decision. Examples include a financial interest
in a supplier, a close personal relationship with a candidate or a direct
report, and outside work with a competitor.

Conflicts must be declared in the HR Portal within five working days of arising.
Declaring a conflict is not itself a problem; failing to declare one is.

An employee must not take part in a decision about a person with whom they have
a close personal relationship, including a hiring, promotion or pay decision.

## 4. Gifts and hospitality

Gifts and hospitality must never be given or accepted in order to obtain an
improper advantage.

Gifts up to a value of 50 units may be accepted without approval. Gifts between
50 and 250 units must be declared in the gifts register and require line manager
approval. Gifts above 250 units must be declined.

Cash and cash equivalents, including vouchers, must always be declined
regardless of value. Hospitality is acceptable where it is reasonable,
infrequent, and connected to a legitimate business discussion.

No gift or hospitality of any value may be offered to a government official
without prior approval from Legal.

## 5. Confidential information

Company, customer and employee information must be protected. Confidential
information must not be discussed in public places, shared with family, or used
for personal benefit. These obligations continue after employment ends.

## 6. Company property and records

Company assets are provided for business use. Business records, including
expenses, timesheets and customer records, must be accurate and complete.
Falsifying a record is gross misconduct.

## 7. Social media and public comment

Employees may say that they work at Nimbus Technologies. They must make clear
that personal views are their own, must not disclose confidential information,
and must not speak on behalf of the company unless authorised.

Media enquiries are directed to the Communications team without comment.

## 8. Anti-bribery, sanctions and competition

We do not offer, give, request or accept bribes or facilitation payments in any
form. We comply with applicable sanctions and export controls.

We compete on merit. Employees must not discuss prices, markets or customers
with a competitor, and must leave any meeting where such a discussion begins,
reporting it to Legal the same day.

## 9. Raising a concern

Concerns may be raised with a line manager, the People team, or the confidential
whistleblowing line, which accepts anonymous reports and is operated by an
independent provider.

Reports are investigated promptly and confidentially. Retaliation against
someone who raises a concern in good faith is itself gross misconduct and will
be treated as such, whether or not the original concern is upheld.

## 10. Gross misconduct

Examples of gross misconduct include theft, fraud, falsification of records,
violence, harassment, being unfit for work through alcohol or drugs, serious
breach of the IT policy, disclosure of confidential information, and retaliation
against a whistleblower. Gross misconduct may result in dismissal without
notice.
"""

COMPANY_FAQS = """
# Company FAQs

Quick answers to the questions the People and IT teams are asked most often.
Where an FAQ conflicts with a policy document, the policy document is correct.

## Pay and payroll

**When am I paid?**
On the 25th of each month, or the preceding working day if the 25th falls on a
weekend or public holiday.

**Where do I find my payslip?**
In the HR Portal, from two working days before payday. Payslips for the last
seven years are available.

**Who do I contact about a payroll error?**
Raise a ticket with the People team through the HR Portal. Payroll corrections
are made in the next available run; an emergency payment can be requested for an
underpayment of more than 500 units.

## Leave

**How many annual leave days do I get?**
25 days for full-time employees, rising to 27 after five years and 30 after ten.
See the Leave Policy for the full detail.

**Can I carry unused leave into next year?**
Up to 5 days, which must be used by 31 March. More than 5 days requires written
approval from your manager and HR before 31 December.

**How much notice do I need to give?**
Twice the length of the leave, with a minimum of five working days. Ten days or
more needs four weeks' notice.

## Working arrangements

**How many days do I need to be in the office?**
Three days a week, and Tuesday is the fixed company-wide office day.

**Can I work from abroad?**
For up to 20 working days a year, from an approved country, with your manager's
approval recorded in the HR Portal. Longer needs People Director and Finance
approval.

**Do I get an allowance for my home office?**
400 units when you join, and 200 units every three years after that.

## IT

**Who do I contact for an IT problem?**
The Service Desk, through the IT portal. For a suspected security incident,
contact them immediately.

**Can I install software myself?**
Only from the Company Portal. Anything else needs a ticket with a business
justification.

**Can I use ChatGPT or another AI tool for work?**
Only company-approved tools, and never with Confidential or Restricted
information. See the IT Acceptable Use Policy.

**I've lost my laptop. What do I do?**
Report it to the Service Desk within 24 hours so it can be wiped remotely.

## Facilities

**Is there parking at the office?**
Limited spaces are allocated by ballot each quarter, with priority for employees
with accessibility needs. Register in the facilities system.

**Do I need an ID card?**
Yes. Your card is issued on your first day and must be worn visibly in the
office. Report a lost card to reception; a replacement costs 15 units.

## Expenses

**How long do I have to claim an expense?**
30 days from the date of expenditure. After 60 days it will not be reimbursed
except in exceptional circumstances.

**Do I need a receipt?**
For anything above 25 units, yes, and it must be itemised. A card statement is
not enough.

**What is the mileage rate?**
0.45 units per kilometre for the first 10,000 kilometres in a year, then 0.25.
"""


PDF_DOCS = [
    ("Leave_Policy.pdf", "Leave Policy", LEAVE_POLICY),
    ("Employee_Handbook.pdf", "Employee Handbook", EMPLOYEE_HANDBOOK),
    ("Benefits_Guide.pdf", "Benefits Guide", BENEFITS_GUIDE),
]

DOCX_DOCS = [
    ("IT_Acceptable_Use_Policy.docx", "IT Acceptable Use Policy", IT_POLICY),
    ("Travel_and_Expenses_Policy.docx", "Travel and Expenses Policy", TRAVEL_POLICY),
    ("Remote_Work_Policy.docx", "Remote Work Policy", REMOTE_WORK_POLICY),
]

MD_DOCS = [
    ("Code_of_Conduct.md", CODE_OF_CONDUCT),
    ("Company_FAQs.md", COMPANY_FAQS),
]


def write_pdf(path: Path, title: str, body: str) -> bool:
    """Write a policy document as a PDF using reportlab."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        return False

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=title)
    story = [
        Paragraph("Nimbus Technologies", styles["Heading3"]),
        Paragraph(title, styles["Title"]),
        Spacer(1, 16),
    ]
    for block in body.strip().split("\n\n"):
        block = block.strip()
        style = "Heading2" if block[0].isdigit() and "\n" not in block else "BodyText"
        story.append(Paragraph(block.replace("\n", " "), styles[style]))
        story.append(Spacer(1, 8))
    doc.build(story)
    return True


def write_docx(path: Path, title: str, body: str) -> None:
    """Write a policy document as a Word file using python-docx."""
    import docx

    document = docx.Document()
    document.add_paragraph("Nimbus Technologies")
    document.add_heading(title, level=0)
    for block in body.strip().split("\n\n"):
        block = block.strip()
        if block[0].isdigit() and "\n" not in block:
            document.add_heading(block, level=1)
        else:
            document.add_paragraph(block.replace("\n", " "))
    document.save(str(path))


def main() -> None:
    """Write every document into the documents/ folder."""
    DOCS.mkdir(parents=True, exist_ok=True)
    written = 0

    for name, title, body in PDF_DOCS:
        if write_pdf(DOCS / name, title, body):
            written += 1
        else:
            print("reportlab is not installed - writing", name, "as .txt instead",
                  file=sys.stderr)
            (DOCS / name.replace(".pdf", ".txt")).write_text(
                f"{title}\n\n{body}", encoding="utf-8"
            )
            written += 1

    for name, title, body in DOCX_DOCS:
        write_docx(DOCS / name, title, body)
        written += 1

    for name, body in MD_DOCS:
        (DOCS / name).write_text(body.strip() + "\n", encoding="utf-8")
        written += 1

    print(f"Wrote {written} documents to {DOCS}")


if __name__ == "__main__":
    main()
